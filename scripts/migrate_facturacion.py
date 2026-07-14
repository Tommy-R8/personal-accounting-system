"""Migrate the Facturación y Cobranzas workbook into Supabase (spec §10.2).

Usage:
  python scripts/migrate_facturacion.py                # dry run
  python scripts/migrate_facturacion.py --execute      # insert + reconcile
"""
import sys
from datetime import datetime

import openpyxl

from lib import db_conn

XLSX = "data/Facturación y Cobranzas.xlsx"


def load_invoices(wb):
    rows, problems = [], []
    for lineno, r in enumerate(wb["Facturas"].iter_rows(min_row=2, values_only=True),
                               start=2):
        if r[0] is None:
            continue
        try:
            pv, num = str(r[3]).split("-")
            rows.append(dict(issued_on=r[0].date(), punto_venta=pv,
                             invoice_number=int(num), amount_ars=float(r[4])))
        except Exception as e:  # noqa: BLE001
            problems.append((lineno, f"Facturas: {e}"))
    return rows, problems


def classify(platform, usd):
    """Kind rule, verified against every row of the sheet (spec §10.2)."""
    if usd is None:
        return "bank_deposit"
    return "cash_sale" if platform == "Cash" else "crypto_sale"


def load_conversions(wb):
    rows, problems = [], []
    for lineno, r in enumerate(
            wb["Cobranzas y Depósitos"].iter_rows(min_row=2, values_only=True),
            start=2):
        if not isinstance(r[0], datetime):
            continue  # trailing formula/blank rows
        ars, platform, usd, tc, notas = r[3], str(r[4]).strip(), r[5], r[6], r[8]
        try:
            usd = float(usd) if isinstance(usd, (int, float)) else None
            kind = classify(platform, usd)
            # rate_quoted only when plausible (the Buenbit row holds the ARS
            # amount in Tipo de Cambio -> import as null)
            quoted = float(tc) if isinstance(tc, (int, float)) and tc < 100000 else None
            rows.append(dict(occurred_on=r[0].date(), kind=kind,
                             ars_amount=float(ars), usd_amount=usd,
                             rate_quoted=quoted, platform=platform,
                             notes=str(notas) if notas else None, line=lineno))
        except Exception as e:  # noqa: BLE001
            problems.append((lineno, str(e)))
    return rows, problems


def metrics(conversions, invoices):
    dep = sum(c["ars_amount"] for c in conversions
              if c["kind"] in ("crypto_sale", "bank_deposit"))
    cob = sum(c["ars_amount"] for c in conversions
              if c["kind"] in ("crypto_sale", "cash_sale"))
    fac = sum(i["amount_ars"] for i in invoices)
    return dict(facturado=fac, depositado=dep, cobrado=cob,
                restante=fac - dep, disponible=cob - dep)


def dry_run(invoices, conversions, problems):
    print(f"\n=== DRY RUN — {len(invoices)} invoices, "
          f"{len(conversions)} conversions, {len(problems)} problems ===")
    for lineno, msg in problems:
        print(f"  line {lineno}: {msg}")
    print("\n-- Conversions classification --")
    for c in conversions:
        print(f"  line {c['line']:3} {c['occurred_on']} {c['kind']:12} "
              f"ARS {c['ars_amount']:>13,.2f} {c['platform']:8} "
              f"USD {c['usd_amount'] or 0:>9,.2f} quoted={c['rate_quoted']}")
    print("\n-- Metrics computed from the rows (compare with the Overview tab) --")
    for k, v in metrics(conversions, invoices).items():
        print(f"  {k:12} {v:>15,.2f}")
    print("\nRe-run with --execute to insert.")


def insert_all(invoices, conversions):
    with db_conn() as conn, conn.cursor() as cur:
        cur.execute("select name, id from payment_methods")
        pm_ids = dict(cur.fetchall())
        cur.executemany(
            """insert into invoices (issued_on, punto_venta, invoice_number,
                                     amount_ars, source)
               values (%(issued_on)s, %(punto_venta)s, %(invoice_number)s,
                       %(amount_ars)s, 'migration')""",
            invoices,
        )
        cur.executemany(
            """insert into conversions (occurred_on, kind, ars_amount, usd_amount,
                                        rate_quoted, platform_id, notes, source)
               values (%(occurred_on)s, %(kind)s, %(ars_amount)s, %(usd_amount)s,
                       %(rate_quoted)s, %(platform_id)s, %(notes)s, 'migration')""",
            [dict(c, platform_id=pm_ids[c["platform"]]) for c in conversions],
        )
        conn.commit()
    print(f"inserted {len(invoices)} invoices + {len(conversions)} conversions")


def check_platforms(conversions):
    """Resolve platform names at dry-run time (not mid-insert): unknown spellings
    in a fresh export become reviewable problems instead of a KeyError."""
    with db_conn() as conn, conn.cursor() as cur:
        cur.execute("select name from payment_methods where active")
        known = {n for (n,) in cur.fetchall()}
    return [(c["line"], f"unknown platform {c['platform']!r} — "
             f"add it to payment_methods or fix the cell")
            for c in conversions if c["platform"] not in known]


def reconcile(invoices, conversions):
    """Spec §10.2: sheet-derived metrics must equal v_deposit_status to the peso.
    Valid because at this point only migrated rows exist in the fiscal tables."""
    with db_conn() as conn, conn.cursor() as cur:
        cur.execute("""select count(*) from invoices where source <> 'migration'
                       union all
                       select count(*) from conversions where source <> 'migration'""")
        foreign = sum(n for (n,) in cur.fetchall())
    if foreign:
        print(f"\nWARNING: {foreign} non-migration rows exist in fiscal tables — "
              "the v_deposit_status comparison below includes them")
    expected = metrics(conversions, invoices)
    with db_conn() as conn, conn.cursor() as cur:
        cur.execute("select * from v_deposit_status")
        cols = [d[0] for d in cur.description]
        actual = dict(zip(cols, (float(v) for v in cur.fetchone())))
    mapping = {  # sheet-side key -> view column
        "facturado": "total_facturado", "depositado": "total_depositado",
        "cobrado": "total_cobrado", "restante": "restante_por_depositar",
        "disponible": "disponible_para_depositar",
    }
    print(f"\n{'metric':12} {'sheet':>16} {'db':>16}")
    ok = True
    for k, col in mapping.items():
        s, d = expected[k], actual[col]
        match = abs(s - d) < 0.01  # to the peso
        ok = ok and match
        print(f"{k:12} {s:>16,.2f} {d:>16,.2f}{'' if match else '  <-- MISMATCH'}")
    print(f"\nRECONCILIATION {'PASS' if ok else 'FAIL'}")


def main():
    wb = openpyxl.load_workbook(XLSX, data_only=True)
    invoices, inv_problems = load_invoices(wb)
    conversions, problems = load_conversions(wb)
    problems = inv_problems + problems
    problems += check_platforms(conversions)
    if "--execute" not in sys.argv:
        dry_run(invoices, conversions, problems)
        return
    if problems:
        sys.exit(f"refusing to execute with {len(problems)} problem rows")
    insert_all(invoices, conversions)
    reconcile(invoices, conversions)


if __name__ == "__main__":
    main()
