"""Migrate the old expense tracker workbook into Supabase (spec §10.1).

Usage:
  python scripts/migrate_from_sheets.py                # dry run
  python scripts/migrate_from_sheets.py --execute      # insert + reconcile
"""
import sys
from collections import defaultdict
from datetime import datetime

import openpyxl

from lib import db_conn

XLSX = "data/Personal Expenses Tracker (responses).xlsx"
SHEET = "Respuestas de formulario 1"

# ---- CONFIG (verified against the workbook on 2026-07-13) --------------------
# 0-based column indices: Marca temporal, Name, Category, Currency, Amount,
#                         ARCA, Blue USD Rate, USD Amount, Month
I_TS, I_NAME, I_CAT, I_CUR, I_AMOUNT, I_ARCA, I_RATE, I_USD, I_MONTH = range(9)
ARCA_NEGRO_VALUES = {"negro"}   # blank/None/'Blanco' -> fiscal=True
CATEGORY_MAP = {                # old sheet name -> new category name (1:1 today)
    "Comida": "Comida", "Salidas": "Salidas", "Transporte": "Transporte",
    "Compras": "Compras", "Regalos": "Regalos", "Suscripciones": "Suscripciones",
    "Impuestos": "Impuestos",
}
COMBO_RULES = {                 # frozenset of NEW names -> winning category.
    # Confirmed by the owner 2026-07-13 during dry-run review (plan Task 10 Step 5).
    frozenset({"Comida", "Salidas"}): "Salidas",
    frozenset({"Compras", "Regalos"}): "Regalos",
    frozenset({"Compras", "Salidas"}): "Salidas",
    frozenset({"Salidas", "Transporte"}): "Salidas",
    frozenset({"Impuestos", "Suscripciones"}): "Suscripciones",
    frozenset({"Regalos", "Salidas"}): "Salidas",
    frozenset({"Comida", "Regalos"}): "Comida",
    frozenset({"Compras", "Transporte"}): "Transporte",
    frozenset({"Comida", "Compras"}): "Comida",
}
LINE_CATEGORY_OVERRIDES = {     # rows with an empty Category cell — confirmed 2026-07-13
    287: "Transporte",   # Uber
    344: "Comida",       # McDonald's
    404: "Salud",        # Farmacia
    462: "Transporte",   # Estacionamiento FIUBA
}
# -------------------------------------------------------------------------------


def resolve_category(raw):
    if not raw or not str(raw).strip():
        raise ValueError("row has no category — assign one manually")
    names = [CATEGORY_MAP[c.strip()] for c in str(raw).split(",") if c.strip()]
    if len(names) == 1:
        return names[0], []
    winner = COMBO_RULES.get(frozenset(names), names[0])
    tags = [n.lower() for n in names if n != winner]
    return winner, tags


def fetch_rates():
    """Backfilled Blue sell rates, for rows whose sheet conversion is broken."""
    with db_conn() as conn, conn.cursor() as cur:
        cur.execute("select rate_date, blue_sell from fx_rates order by rate_date")
        return cur.fetchall()  # ascending [(date, rate), ...]


def sell_rate_on(rates, day):
    """Latest rate on or before `day` (same fallback rule as add_transaction)."""
    best = None
    for rate_date, rate in rates:
        if rate_date > day:
            break
        best = rate
    if best is None:
        raise ValueError(f"no Blue rate on or before {day}")
    return float(best)


def load_rows():
    rates = fetch_rates()
    ws = openpyxl.load_workbook(XLSX, read_only=True, data_only=True)[SHEET]
    rows, problems = [], []
    for lineno, r in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
        if r[I_TS] is None:
            continue  # trailing blank rows
        try:
            occurred_on = r[I_TS].date()
            amount = float(r[I_AMOUNT])
            currency = (str(r[I_CUR] or "ARS")).strip().upper()
            if lineno in LINE_CATEGORY_OVERRIDES:
                category, tags = LINE_CATEGORY_OVERRIDES[lineno], []
            else:
                category, tags = resolve_category(r[I_CAT])
            arca = str(r[I_ARCA] or "").strip().lower()
            fiscal = arca not in ARCA_NEGRO_VALUES
            repaired = False
            if currency == "USD":
                usd, usd_rate = amount, 1.0
            elif isinstance(r[I_USD], (int, float)):
                usd = float(r[I_USD])           # the sheet's own conversion
                usd_rate = amount / usd
            else:
                # 19 known rows: Bluelytics 429 error text / #VALUE! in the sheet.
                usd_rate = sell_rate_on(rates, occurred_on)
                usd = amount / usd_rate
                repaired = True
            rows.append(dict(
                occurred_on=occurred_on,
                description=str(r[I_NAME] or "").strip() or "(sin descripción)",
                category=category, tags=tags, amount=amount, currency=currency,
                fiscal=fiscal, usd=usd, usd_rate=usd_rate, repaired=repaired,
                line=lineno,
            ))
        except Exception as e:  # noqa: BLE001 — report every bad row, don't die
            problems.append((lineno, f"{r[I_NAME]!r}: {e}"))
    return rows, problems


def dry_run_report(rows, problems):
    multi = [r for r in rows if r["tags"]]
    repaired = [r for r in rows if r["repaired"]]
    print(f"\n=== DRY RUN — {len(rows)} importable rows, {len(problems)} problems ===")
    if problems:
        print("\n-- Problem rows (fix in workbook or CONFIG, then rerun) --")
        for lineno, msg in problems:
            print(f"  line {lineno}: {msg}")
    if multi:
        print(f"\n-- Multi-category rows ({len(multi)}) — winner + tags --")
        for r in multi:
            print(f"  line {r['line']:5} {r['occurred_on']} "
                  f"{r['description'][:40]:40} -> {r['category']} tags={r['tags']}")
    if repaired:
        print(f"\n-- Repaired-rate rows ({len(repaired)}) — sheet conversion was "
              "broken; recomputed from fx_rates --")
        for r in repaired:
            print(f"  line {r['line']:5} {r['occurred_on']} "
                  f"{r['description'][:40]:40} rate={r['usd_rate']:.2f}")
    per_month = defaultdict(float)
    for r in rows:
        per_month[r["occurred_on"].strftime("%Y-%m")] += r["usd"]
    print("\n-- Per-month USD totals --")
    for m in sorted(per_month):
        print(f"  {m}  {per_month[m]:12.2f}")
    print("\nRe-run with --execute to insert.")


def insert_rows(rows):
    with db_conn() as conn, conn.cursor() as cur:
        cur.execute("select count(*) from transactions where source = 'migration'")
        existing = cur.fetchone()[0]
        if existing:
            sys.exit(f"refusing: {existing} source='migration' rows already exist — "
                     "delete them first (see plan Task 10 Step 6) and re-run")
        cur.execute("select name, id from categories")
        cat_ids = dict(cur.fetchall())
        cur.executemany(
            """insert into transactions
                 (occurred_on, type, description, category_id, amount, currency,
                  fiscal, tags, usd_rate, source)
               values (%(occurred_on)s, 'outflow', %(description)s, %(category_id)s,
                       %(amount)s, %(currency)s, %(fiscal)s, %(tags)s,
                       %(usd_rate)s, 'migration')""",
            [dict(r, category_id=cat_ids[r["category"]]) for r in rows],
        )
        conn.commit()
    print(f"inserted {len(rows)} rows with source='migration'")


def sheet_side():
    """Independent raw pass over the worksheet — deliberately NOT reusing the
    parsing pipeline, so a row load_rows() drops or mangles cannot cancel out
    of the comparison. Sums only the sheet's own numeric USD values."""
    ws = openpyxl.load_workbook(XLSX, read_only=True, data_only=True)[SHEET]
    per_month, n_rows, n_broken = defaultdict(float), 0, 0
    for r in ws.iter_rows(min_row=2, values_only=True):
        if r[I_TS] is None:
            continue
        n_rows += 1
        if isinstance(r[I_USD], (int, float)):
            per_month[r[I_TS].strftime("%Y-%m")] += float(r[I_USD])
        else:
            n_broken += 1  # broken conversions: repaired rows, excluded here
    return per_month, n_rows, n_broken


def reconcile(rows):
    per_month_sheet, n_sheet_rows, n_broken = sheet_side()
    with db_conn() as conn, conn.cursor() as cur:
        cur.execute("select count(*) from transactions where source = 'migration'")
        n_db = cur.fetchone()[0]
        cur.execute(
            """select to_char(t.occurred_on, 'YYYY-MM'), sum(t.amount_usd)
                 from transactions t
                where t.source = 'migration'
                group by 1"""
        )
        per_month_db = {m: float(t) for m, t in cur.fetchall()}
    # db totals include repaired rows; subtract their recomputed USD per month
    # so both sides cover the same clean-row population (spec §10.1).
    repaired_by_month = defaultdict(float)
    for r in rows:
        if r["repaired"]:
            repaired_by_month[r["occurred_on"].strftime("%Y-%m")] += r["usd"]
    print(f"\nrow counts: sheet={n_sheet_rows} db={n_db} "
          f"(sheet rows with broken USD: {n_broken})")
    print(f"{'month':8} {'sheet_usd':>12} {'db_usd':>12} {'diff%':>8}")
    worst = 0.0
    for m in sorted(per_month_sheet):
        s = per_month_sheet[m]
        d = per_month_db.get(m, 0.0) - repaired_by_month.get(m, 0.0)
        diff = abs(s - d) / s * 100 if s else 0.0
        worst = max(worst, diff)
        print(f"{m:8} {s:12.2f} {d:12.2f} {diff:8.3f}")
    counts_ok = n_db == n_sheet_rows
    if not counts_ok:
        print(f"\nROW COUNT MISMATCH: sheet has {n_sheet_rows}, db has {n_db}")
    verdict = "PASS" if worst < 0.5 and counts_ok else "FAIL"
    print(f"\nRECONCILIATION {verdict} — worst monthly diff {worst:.3f}% (limit 0.5%)")


def main():
    rows, problems = load_rows()
    if "--execute" not in sys.argv:
        dry_run_report(rows, problems)
        return
    if problems:
        sys.exit(f"refusing to execute with {len(problems)} problem rows — "
                 "run the dry run and fix them first")
    insert_rows(rows)
    reconcile(rows)


if __name__ == "__main__":
    main()
